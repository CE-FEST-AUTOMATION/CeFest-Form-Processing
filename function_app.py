import azure.functions as func
import logging
import json
import urllib.request
import pandas as pd
import numpy as np
import openpyxl
import re

# Create the FunctionApp instance used by the Functions host
app = func.FunctionApp()


#endpoint for uploading 
ENDPOINT = "https://inpro-participants-submit-3avzbu5vaa-as.a.run.app"

#pathing to file inside onedrive locally:
path = "C:\\Users\\xuan2\\Downloads\\responses.xlsx"

#mapping to change according to the form (change left side)
mapping = { "tele handle": "tele_handle",
           "full name": "full_name",
            "Student Email": "student_email",
            "First Choice": "trail_first_choice",
            "Second Choice": "trail_second_choice",
            "Third Choice": "trail_third_choice"}

#Functions to help with counter
def load_counter(): #counter needs to be saved to cell "A1" of the second sheet labelled "counter"
    wb = openpyxl.load_workbook(path)
    if "counter" not in wb.sheetnames:
        ws = wb.create_sheet("counter")
        ws["A1"] = 0
        wb.save(path)
        return 0
    val = wb["counter"]["A1"].value
    return int(val) if val is not None else 0

def save_counter(value: int):
    wb = openpyxl.load_workbook(path)
    if "counter" not in wb.sheetnames:
        wb.create_sheet("counter")
    wb["counter"]["A1"] = value
    wb.save(path)

counter = load_counter() #counter to keep track of number of rows uploaded

#FUNCTIONS TO CONVERT 1 ROW INTO PAYLOAD
def parse_numeric_from_string(v):
    # return None for NaN-like values
    if pd.isna(v):
        return None
    if isinstance(v, (int, float)):
        try:
            return int(v)
        except Exception:
            pass
    s = str(v)
    m = re.search(r'(\d+)', s)
    return int(m.group(1)) if m else s

def format_vals(init: dict) -> dict:
    if 'year' in init:
        init['year'] = parse_numeric_from_string(init['year'])
    for k in ('trail_first_choice', 'trail_second_choice', 'trail_third_choice'):
        if k in init:
            init[k] = parse_numeric_from_string(init[k])
    return init

def normalize_value(v):
    if pd.isna(v):
        return None
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.floating,)):
        return float(v)
    if isinstance(v, (np.bool_,)):
        return bool(v)
    return v

def process_row(srs: pd.Series) -> dict:
    """Accepts a pandas Series (one row) and returns a JSON-serializable dict."""
    cols = ['tele_handle', 'full_name', 'house', 'student_email', 'year', 'status', 'role', 'trail_first_choice', 'trail_second_choice', 'trail_third_choice'] #edit col names according to the final one
    d = {}
    for col_name in cols:
        if col_name == 'role':
            val = srs.get('role')
            d['role'] = val if not pd.isna(val) else "Participant"
        elif col_name == 'status':
            val = srs.get('status')
            d['status'] = val if not pd.isna(val) else 'pending'
        else:
            d[col_name] = srs.get(col_name)
    d = format_vals(d)
    # normalize numpy / pandas types to Python builtins for JSON
    for k in list(d.keys()):
        d[k] = normalize_value(d[k])
    return d


@app.timer_trigger(schedule="0 * * * * *", arg_name="myTimer", run_on_startup=False)
def check_forms(myTimer: func.TimerRequest) -> None:
    global counter
    # logging.info("Python timer trigger function executed.")
    if myTimer.past_due: # for tracking if the timer is due on time  
        logging.info("The timer is past due.")
    df = pd.read_excel(path, sheet_name="response").rename(columns=mapping)
    curr_num = len(df)
    if counter >= curr_num:
        return
    backlog = df.iloc[counter:curr_num]
    for _, row in backlog.iterrows():
        payload = process_row(row)
        upload_data(payload)
        counter += 1
        save_counter(counter)


def upload_data(payload : dict): #function to upload data onto firebase (DONT NEED TOUCH)
    data = json.dumps(payload).encode()
    req = urllib.request.Request(
        ENDPOINT,
        data=data,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req) as resp:
            body = resp.read().decode()
            print(f"Status: {resp.status}")
            print(f"Response: {body}")
    except urllib.error.HTTPError as e:
        body = e.read().decode()
        print(f"Status: {e.code}")
        print(f"Error: {body}")